"""XLSX: спецификация строками таблицы, первая строка — заголовки."""
import io

from openpyxl import load_workbook

from app.parsers import Block, ParseError


def parse_xlsx(data: bytes) -> list[Block]:
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise ParseError(f"не удалось открыть XLSX: {exc}") from exc

    blocks: list[Block] = []
    for number, sheet in enumerate(workbook.worksheets, start=1):
        rows = [
            ["" if cell is None else str(cell).strip() for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        rows = [row for row in rows if any(row)]
        if len(rows) < 2:
            continue
        if not any(rows[0]):
            raise ParseError(f"лист «{sheet.title}» без заголовков колонок")
        blocks.append(
            Block(
                content="\n".join(" | ".join(row) for row in rows),
                kind="table",
                rows=rows,
                metadata={"page": number, "section": sheet.title},
            )
        )

    if not blocks:
        raise ParseError("в XLSX нет строк со значениями")
    return blocks
